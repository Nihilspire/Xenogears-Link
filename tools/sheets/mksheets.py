"""Hint sheet and room list for playtesting (20 Oct 2026). Data: /tmp/claude-0/appdata.json (the app's own tables, exported by node).
Hint sheet: Jack's last sheet (1.99.0.22) rewritten from hints.js, keeping his Checked and notes columns. Room list: one row per room."""
import json,sys,copy
from openpyxl import load_workbook,Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
VER=sys.argv[1]
D=json.load(open('/tmp/claude-0/appdata.json'))
ROOMS={int(k):v for k,v in D['ROOMS'].items()};AR=D['AREAS'];CH=D['CHAPTERS']
def area(f):
    if f<0: return 'World map'
    for r in AR:
        if r[0]<=f<=r[1]: return r[2]
    return '(not named yet)'
def rname(f): return 'World map' if f<0 else str(f)+' '+(ROOMS.get(f) or '(no name)')
def chap(k):
    c=None
    for s,n in CH:
        if s<=k: c=n
    return c or ''
def rooms_for(k):
    k=str(k)
    if k in D['HPATH']: L=[(q[0],q[1]) for q in D['HPATH'][k]]
    else:
        pre=D['HPRE'].get(k,[]);rows=D['HINTS'].get(k,[])
        L=[(q[0],q[1]) for q in pre]+[(r[0],[[r[2],r[3]]] if len(r)>3 else []) for r in rows[:max(1,3-len(pre))]]
    return L
def fmt(L): return ', '.join(('World' if f<0 else str(f))+(' ('+'; '.join('%d, %d'%(q[0],q[1]) for q in sp)+')' if sp else '') for f,sp in L)
def parts(t):
    import re
    if not t: return []
    t=t[1:] if t.startswith('?') else t
    return [x.strip() for x in re.split(r'(?<=\.)\s+',t) if x.strip()]
# ---------- hint sheet ----------
src='/mnt/user-data/outputs/Xenogears story hints 1.99.0.22.xlsx'
wb=load_workbook(src);ws=wb.active
ws['A1']='Xenogears Link %s - story hints'%VER
hdr=[c.value for c in ws[3]];col={h:i+1 for i,h in enumerate(hdr)}
ws.cell(3,col['Changed since 1.85']).value='Changed since 1.99.0.22'
rows={}
for r in range(4,ws.max_row+1):
    v=ws.cell(r,1).value
    if v is not None: rows[int(v)]=r
steps=sorted(set(map(int,D['HINTS']))|set(map(int,D['HNOTE'])))
mism=[]
for k in steps:
    if k not in rows: continue
    r=rows[k];L=rooms_for(k);note=D['HNOTE'].get(str(k),'');note=note[1:] if note.startswith('?') else note   # the sheet keeps '?' as Checked = N
    newE=fmt(L);changed=[]
    if (ws.cell(r,col['Rooms in order, with spots (edit)']).value or '')!=newE: changed.append('rooms')
    if (ws.cell(r,col['Hint text (edit)']).value or '')!=note: changed.append('text')
    ws.cell(r,col['Rooms in order, with spots (edit)']).value=newE
    ws.cell(r,col['Hint text (edit)']).value=note
    ws.cell(r,col['Area']).value=area(L[-1][0]) if L else ''
    ws.cell(r,col['Changed since 1.85']).value=', '.join(changed) or None
    ws.cell(r,col['Room names (in order)']).value='\n'.join(rname(f) for f,sp in L)
    np_,nr=len(parts(note)),len(L)
    if np_ and nr and np_!=nr: mism.append((k,np_,nr))
out='/mnt/user-data/outputs/Xenogears story hints %s.xlsx'%VER;wb.save(out);print('hints',len(steps),'mismatch',mism)
json.dump(mism,open('/tmp/claude-0/mism.json','w'))
# ---------- room list ----------
H=Font(name='Arial',bold=True,color='FFFFFF');F=Font(name='Arial',size=10);fill=PatternFill('solid',fgColor='1F3A5F');ed=PatternFill('solid',fgColor='FFF2CC')
thin=Side(style='thin',color='BBBBBB');B=Border(left=thin,right=thin,top=thin,bottom=thin)
hintRooms={}
for k in steps:
    for f,sp in rooms_for(k): hintRooms.setdefault(f,[]).append(k)
miss={};[miss.setdefault(int(k.split(':')[0]),0) for k in D['MISS']['tag']]
for k in D['MISS']['tag']: miss[int(k.split(':')[0])]+=1
for r,L in D['MISS']['extra'].items(): miss[int(r)]=miss.get(int(r),0)+len(L)
wb2=Workbook();ws2=wb2.active;ws2.title='Rooms'
ws2['A1']='Xenogears Link %s - room list'%VER;ws2['A1'].font=Font(name='Arial',bold=True,size=14)
ws2['A2']='Every room with a name or a use. Yellow columns are yours: mark a room once you have walked it with the app, and note anything off (doors, spots, names).';ws2['A2'].font=Font(name='Arial',italic=True,size=10)
cols=['Room','Area','Name','Ways out (to rooms)','Doors without a spot','Treasure / gifts','Spots missing','Missables','Saves / shops','Used by hint steps','Walked (Y/N)','Notes']
W=[7,20,30,34,10,46,9,9,14,22,10,40]
for i,(c,w) in enumerate(zip(cols,W)):
    x=ws2.cell(3,i+1,c);x.font=H;x.fill=fill;x.alignment=Alignment(wrap_text=True,vertical='center');x.border=B
    ws2.column_dimensions[chr(65+i)].width=w
MK={'s':'save','i':'item shop','g':'gear shop'}
r=4
allrooms=sorted(set(ROOMS)|set(int(k) for k in D['EXITS'])|set(int(k) for k in D['TRE']))
for f in allrooms:
    nmv=ROOMS.get(f,'')
    if nmv in('Cutscene','???','Debug','Debug Map') and str(f) not in D['EXITS'] and str(f) not in D['TRE']: continue
    E=D['EXITS'].get(str(f),[]);T=D['TRE'].get(str(f),[]);M=D['MARKS'].get(str(f),[])
    outs=sorted(set(('W' if e[2]<0 else str(e[2])) for e in E),key=lambda x:(x=='W',int(x) if x!='W' else 0))
    vals=[f,area(f),nmv or '(no name)',', '.join(outs),sum(1 for e in E if e[0] is None and e[2]>=0) or None,
          '; '.join(t['n']+(' (gift)' if t['gift'] else '') for t in T) or None,sum(1 for t in T if t['x'] is None) or None,miss.get(f) or None,
          ', '.join(sorted(set(MK.get(m[2],m[2]) for m in M))) or None,', '.join(map(str,hintRooms.get(f,[]))) or None,None,None]
    for i,v in enumerate(vals):
        x=ws2.cell(r,i+1,v);x.font=F;x.border=B;x.alignment=Alignment(wrap_text=True,vertical='top')
        if i>=10: x.fill=ed
    r+=1
ws2.freeze_panes='D4';ws2.auto_filter.ref='A3:L%d'%(r-1)
dv=DataValidation(type='list',formula1='"Y,N"',allow_blank=True);ws2.add_data_validation(dv);dv.add('K4:K%d'%(r-1))
out2='/mnt/user-data/outputs/Xenogears room list %s.xlsx'%VER;wb2.save(out2);print('rooms',r-4)
