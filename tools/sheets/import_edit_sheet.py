#!/usr/bin/env python3
"""Loads Jack's "story hints edit sheet" (.xlsx, made by tools/sheets/mk_edit_sheet.py) back into app/www/hints.js and hsplit.js.
Rows with the same Step (and Part) are one hint's pages in order; each row is one page: X, Z is its green spot and More spots holds any others (x, z; x, z).
Part 1 (or blank) is the step's own note and path; Part 2+ are the later parts in HSPLIT (only their text and path change; the switches
that start them stay). "In game" Yes = checked (no "?" in front).   usage: import_edit_sheet.py sheet.xlsx [--dry]"""
import sys,json,re,os
from openpyxl import load_workbook
src=sys.argv[1];dry='--dry' in sys.argv;W=os.path.join(os.path.dirname(__file__),'../../app/www/')
ws=load_workbook(src,data_only=True)['Hints'];H={}
for r in ws.iter_rows(min_row=2,values_only=True):
    st,pt,tx,room,_,_,x,z,more,ing=r[:10]
    if st is None or str(st).strip()=='': continue
    k=int(st);p=int(pt) if pt not in (None,'') else 1;e=H.setdefault((k,p),{'text':None,'ok':False,'path':[]})
    if tx not in (None,'') and e['text'] is None: e['text']=str(tx).strip();e['ok']=str(ing or '').strip().lower() in ('yes','y','✓','x')
    if room in (None,''): continue
    room=int(room);sp=[[int(round(float(x))),int(round(float(z)))]] if x not in (None,'') and z not in (None,'') else []
    for m in re.findall(r'(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)',str(more or '')): sp.append([int(round(float(m[0]))),int(round(float(m[1])))])
    e['path'].append([room,sp])
txt=open(W+'hints.js').read();hs=open(W+'hsplit.js').read()
get=lambda t,n:json.loads(re.search(r'^window\.'+n+r'=(.*);\s*$',t,re.M).group(1))
N=get(txt,'HNOTE');P=get(txt,'HPATH');S=json.loads(re.search(r'window\.HSPLIT=(\{.*\});',hs).group(1));ch=0
for (k,p),e in sorted(H.items()):
    note=(('' if e['ok'] else '?')+e['text']) if e['text'] else None;ks=str(k)
    if p==1:
        was=json.dumps([N.get(ks),P.get(ks)])
        if note: N[ks]=note
        if e['path']: P[ks]=e['path']
        if json.dumps([N.get(ks),P.get(ks)])!=was: ch+=1;print('step',k,'changed')
    else:
        sp=S.get(ks) or [];i=p-2
        if i>=len(sp): print('step',k,'part',p,': no such part in the game data, skipped (ask Claude to add the switch for it)');continue
        was=json.dumps(sp[i]);
        while len(sp[i])<5: sp[i].append(None)
        if note: sp[i][3]=note
        if e['path']: sp[i][4]=e['path']
        if json.dumps(sp[i])!=was: ch+=1;print('step',k,'part',p,'changed')
print(ch,'changed'+(' (dry run)' if dry else ''))
if not dry and ch:
    put=lambda t,n,O:re.sub(r'^window\.'+n+r'=.*;\s*$',lambda m:'window.'+n+'={'+','.join(json.dumps(k)+':'+json.dumps(O[k]) for k in sorted(O,key=int))+'};',t,flags=re.M)
    open(W+'hints.js','w').write(put(put(txt,'HNOTE',N),'HPATH',P))
    open(W+'hsplit.js','w').write(hs.replace(re.search(r'window\.HSPLIT=(\{.*\});',hs).group(0),'window.HSPLIT='+json.dumps(S,separators=(',',':'))+';'))
