import json
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,PatternFill,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
D=json.load(open('/tmp/claude-0/hsheet.json'));ROOMS=json.load(open('/tmp/claude-0/rooms_list.json'))
wb=Workbook();ws=wb.active;ws.title='Hints'
F='Arial';hdr=['Step','Part','Hint text','Room ID','Room name','Area','X','Z','More spots','In game','Checked','Notes']
W=[7,6,60,9,30,16,9,9,16,9,9,30]
ws.append(hdr)
for i,w in enumerate(W):ws.column_dimensions[chr(65+i)].width=w
hf=PatternFill('solid',start_color='1F2A3A');thin=Side(style='thin',color='BBBBBB');top=Side(style='medium',color='555555')
for c in ws[1]:c.font=Font(name=F,bold=True,color='FFFFFF');c.fill=hf;c.alignment=Alignment(vertical='center',wrap_text=True)
edit=PatternFill('solid',start_color='FFF8DC')
r=2;band=False
for h in D:
    rows=[]
    for q in (h['rooms'] or [{'f':None,'spots':[]}]):
        rows.append((q['f'],q['spots']))
    band=not band
    for j,(f,s) in enumerate(rows):
        ws.cell(r,1,h['k']);ws.cell(r,2,h['part'] or None)
        if j==0: ws.cell(r,3,h['text'])
        if f is not None: ws.cell(r,4,f)
        ws.cell(r,5,'=IF(D%d="","",IFERROR(VLOOKUP(D%d,Rooms!$A:$B,2,FALSE)&"","?"))'%(r,r))
        ws.cell(r,6,'=IF(D%d="","",IFERROR(VLOOKUP(D%d,Rooms!$A:$C,3,FALSE)&"",""))'%(r,r))
        if s: ws.cell(r,7,s[0][0]);ws.cell(r,8,s[0][1])
        if s and len(s)>1: ws.cell(r,9,'; '.join('%d, %d'%(a,b) for a,b in s[1:]))
        if j==0: ws.cell(r,10,'Yes' if h['chk'] else None)
        for c in range(1,13):
            cell=ws.cell(r,c);cell.font=Font(name=F,size=10,bold=(c==1));cell.alignment=Alignment(vertical='top',wrap_text=(c in (3,5,9,12)))
            cell.border=Border(top=top if j==0 else None,bottom=thin)
            if c in (3,4,7,8,9,11,12):cell.fill=edit
        r+=1
last=r+200
for rr in range(r,last):
    ws.cell(rr,5,'=IF(D%d="","",IFERROR(VLOOKUP(D%d,Rooms!$A:$B,2,FALSE)&"","?"))'%(rr,rr));ws.cell(rr,6,'=IF(D%d="","",IFERROR(VLOOKUP(D%d,Rooms!$A:$C,3,FALSE)&"",""))'%(rr,rr))
    for c in range(1,13):ws.cell(rr,c).font=Font(name=F,size=10)
dv=DataValidation(type='list',formula1='"✓"',allow_blank=True);ws.add_data_validation(dv);dv.add('K2:K%d'%last)
dv2=DataValidation(type='list',formula1='"Yes"',allow_blank=True);ws.add_data_validation(dv2);dv2.add('J2:J%d'%last)
ws.freeze_panes='D2';ws.auto_filter.ref='A1:L%d'%(r-1)
ws.print_title_rows='1:1';ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
ws.print_options.gridLines=True;ws.page_margins.left=ws.page_margins.right=0.3
ws['C1'].comment=Comment('Edit freely. Text goes on the first row of a step; the rows under it with the same Step (and Part) are that hint\'s pages, in order.','Claude')
ws['D1'].comment=Comment('Room number; -1 = world map. Room name fills in by itself.','Claude')
ws['G1'].comment=Comment('Green spot for this page. Extra spots in the same page go in More spots as x, z; x, z','Claude')
# how-to sheet
h=wb.create_sheet('How to use',0)
L=['Xenogears Link story hints (2.0.0.21) — edit sheet','',
'Each row is one page of a hint: a room, and its green spot (X, Z) if it has one. Rows with the same Step (and Part) belong to one hint, in order.',
'Hint text: write it on the first row of the step. Leave it blank on the rows under it.',
'Add a room: insert a row under the step, type the Step (and Part), the Room ID and the spot. The room name fills in from the Rooms tab. -1 is the world map.',
'Two spots on one page: put the first in X, Z and the others in More spots as  x, z; x, z. Two rows in the same room are two separate pages.',
'Add a new hint: add rows with a new Step number.',
'In game: "Yes" = the hint shows in white in the app (checked in game). Leave blank for not checked.',
'Checked: your own tick (✓) while going through the sheet. Notes: anything for me.',
'Cream cells are the ones to edit. Part numbers: hints that change partway through the step (like 114, the Kislev sewers) have Part 1, 2, 3.',
'When done, send the file back and I will load it into the app and the room viewer.',
'Example row: Step 10 | Part (blank) | Climb the Mountain Path. Talk to Yui. | 15 | (Mountain Path) | | -232 | 2190 | | Yes']
for i,t in enumerate(L,1):
    c=h.cell(i,1,t);c.font=Font(name=F,size=13 if i==1 else 10,bold=i==1);c.alignment=Alignment(wrap_text=True,vertical='top')
h.column_dimensions['A'].width=120
rs=wb.create_sheet('Rooms');rs.append(['Room ID','Room name','Area'])
for c in rs[1]:c.font=Font(name=F,bold=True)
rs.append([-1,'World map',''])
for f,n,a in ROOMS: rs.append([f,n,a])
for row in rs.iter_rows(min_row=2):
    for c in row:c.font=Font(name=F,size=10)
rs.column_dimensions['B'].width=40;rs.column_dimensions['C'].width=24
wb.active=1
wb.save('/mnt/user-data/outputs/Xenogears story hints edit sheet.xlsx');print(r-2,'rows')
