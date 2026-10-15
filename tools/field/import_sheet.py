import json,re,sys
from openpyxl import load_workbook
SRC=sys.argv[1]
ws=load_workbook(SRC).active
def split_top(s):
    out=[];d=0;cur=''
    for ch in s:
        if ch=='(':d+=1
        if ch==')':d-=1
        if ch==',' and d==0: out.append(cur);cur=''
        else: cur+=ch
    out.append(cur); return [t.strip().rstrip('.').strip() for t in out if t.strip().rstrip('.').strip()]
def parse(cell):
    if cell is None: return []
    res=[]
    for t in split_top(str(cell)):
        m=re.match(r'^(World|\d+)\s*(?:\((.*)\))?$',t,re.I)
        if not m: raise ValueError('cannot read %r'%t)
        room=-1 if m.group(1).lower()=='world' else int(m.group(1))
        spots=[]
        if m.group(2) and m.group(2).strip():
            for pr in m.group(2).split(';'):
                a=[int(v) for v in re.findall(r'-?\d+',pr)]
                if len(a)==2: spots.append(a)
                elif a: raise ValueError('bad spot %r in %r'%(pr,t))
        res.append([room,spots])
    return res
notes={};rooms={};spots={}
for r in ws.iter_rows(min_row=4,values_only=True):
    if r[0] is None: continue
    s=str(int(r[0])); txt=(r[5] or '').strip(); chk=(r[6] or '').strip().upper()=='Y'
    if txt: notes[s]=('' if chk else '?')+txt
    L=parse(r[4])
    if L: rooms[s]=[q[0] for q in L]; spots[s]=L
json.dump(notes,open('hint_notes.json','w'),indent=1,ensure_ascii=False)
json.dump(rooms,open('hint_rooms.json','w'),indent=1)
json.dump(spots,open('hint_spots.json','w'),indent=0)
print(len(notes),'notes',len(rooms),'room lists',sum(1 for L in spots.values() for q in L if q[0]==-1),'world entries')
for s in ('27','38','42','81','111','114'): print(s,spots.get(s),notes.get(s,'')[:200])
